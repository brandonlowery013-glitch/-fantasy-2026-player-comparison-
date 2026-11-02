"""Fantasy 2026: frozen forecasts, final-result grading and weekly calibration.

Run with the bundled Python runtime: learning.py cycle [--offline].
The spreadsheet is written only by the companion Artifact Tool renderer.
"""
from __future__ import annotations
import argparse, csv, datetime as dt, fcntl, hashlib, io, json, math, os
from pathlib import Path
import sqlite3, statistics, subprocess, urllib.request
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
STATE = Path(os.environ.get('CTD_LEARNING_STATE',str(ROOT.parent / 'learning-state')))
WORKSPACE = ROOT.parents[2]
CACHE = STATE / 'sources'
UTC = dt.timezone.utc
METRICS = ['pass_yards','rush_yards','receptions','receiving_yards','pass_tds',
           'rush_tds','receiving_tds','interceptions','fumbles','ppr']
MAPPING = {'attempts':'pass_attempts','passing_yards':'pass_yards',
 'passing_tds':'pass_tds','passing_interceptions':'interceptions',
 'carries':'rush_attempts','rushing_yards':'rush_yards','rushing_tds':'rush_tds',
 'fantasy_points_ppr':'ppr','fumbles_lost_total':'fumbles'}
URLS = {
 'stats.csv':'https://github.com/nflverse/nflverse-data/releases/download/stats_player/stats_player_week_2026.csv',
 'roster.csv':'https://github.com/nflverse/nflverse-data/releases/download/rosters/roster_2026.csv',
 'snaps.csv':'https://github.com/nflverse/nflverse-data/releases/download/snap_counts/snap_counts_2026.csv',
 'injuries.csv':'https://github.com/nflverse/nflverse-data/releases/download/injuries/injuries_2026.csv',
 'games.csv':'https://raw.githubusercontent.com/nflverse/nfldata/master/data/games.csv',
 'espn.json':'https://site.api.espn.com/apis/site/v2/sports/football/nfl/injuries'}
NODE='/Users/brandonlowery/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'

def iso(t): return t.astimezone(UTC).isoformat()
def utcnow():return dt.datetime.now(UTC)
def parse(s): return dt.datetime.fromisoformat(str(s).replace('Z','+00:00')).astimezone(UTC)
def norm(s):
 import re
 return re.sub('[^a-z0-9]','',str(s).lower().replace(' iii','').replace(' ii','').replace(' jr.','').replace(' jr','').replace(' sr.',''))
def num(v):
 try: return float(v) if v not in ('',None) else 0.
 except (TypeError,ValueError): return 0.
def digest(obj): return hashlib.sha256(json.dumps(obj,sort_keys=True,separators=(',',':'),allow_nan=False).encode()).hexdigest()
def readcsv(p): return list(csv.DictReader(Path(p).open()))
def save_json(path,obj):
 path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
 tmp=path.with_suffix(path.suffix+'.tmp');tmp.write_text(json.dumps(obj,indent=2,allow_nan=False));os.replace(tmp,path)
def fetch(url,path):
 req=urllib.request.Request(url,headers={'User-Agent':'Fantasy2026-research/1.0'})
 with urllib.request.urlopen(req,timeout=45) as response: data=response.read()
 p=Path(path);p.parent.mkdir(parents=True,exist_ok=True);tmp=p.with_suffix(p.suffix+'.tmp');tmp.write_bytes(data);os.replace(tmp,p)
 return {'url':url,'observed_at':iso(dt.datetime.now(UTC)),'sha256':hashlib.sha256(data).hexdigest()}
def db_open(state=STATE):
 state.mkdir(parents=True,exist_ok=True);db=sqlite3.connect(state/'learning.sqlite3')
 db.executescript('''
 CREATE TABLE IF NOT EXISTS forecasts(id TEXT PRIMARY KEY, target TEXT NOT NULL,
   season INTEGER, week INTEGER, captured_at TEXT, kickoff TEXT, payload TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS grades(id TEXT PRIMARY KEY, target TEXT NOT NULL,
   forecast_id TEXT NOT NULL REFERENCES forecasts(id), graded_at TEXT, payload TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS parameter_versions(id TEXT PRIMARY KEY, created_at TEXT, payload TEXT NOT NULL);
 CREATE TRIGGER IF NOT EXISTS forecasts_no_update BEFORE UPDATE ON forecasts BEGIN SELECT RAISE(ABORT,'frozen forecast'); END;
 CREATE TRIGGER IF NOT EXISTS forecasts_no_delete BEFORE DELETE ON forecasts BEGIN SELECT RAISE(ABORT,'frozen forecast'); END;
 CREATE TRIGGER IF NOT EXISTS grades_no_update BEFORE UPDATE ON grades BEGIN SELECT RAISE(ABORT,'append-only grades'); END;
 CREATE TRIGGER IF NOT EXISTS grades_no_delete BEFORE DELETE ON grades BEGIN SELECT RAISE(ABORT,'append-only grades'); END;
 CREATE TRIGGER IF NOT EXISTS parameters_no_update BEFORE UPDATE ON parameter_versions BEGIN SELECT RAISE(ABORT,'versioned parameters'); END;
 ''');return db
def freeze(db,record,state=STATE):
 """Refuse backdating. Caller cannot supply captured_at or promote an old forecast."""
 if 'captured_at' in record: raise ValueError('Caller may not set capture time')
 now=utcnow()
 if parse(record['kickoff'])<=now:return False
 if parse(record['inputs_observed_at'])>now:raise ValueError('Future source timestamp')
 if now-parse(record['inputs_observed_at'])>dt.timedelta(hours=24):raise ValueError('Stale source snapshot')
 if record['data_through_week']>=record['week']:raise ValueError('Target-week leakage')
 r=dict(record,captured_at=iso(now));r['id']=digest(r)
 target=f"{r['season']}:{r['week']}:{r['game_id']}:{r['player_id']}:{r['metric']}"
 path=state/'forecasts'/f"{r['id']}.json";path.parent.mkdir(parents=True,exist_ok=True)
 if not path.exists():
  with path.open('x') as f:json.dump(r,f,indent=2,allow_nan=False)
 db.execute('INSERT OR IGNORE INTO forecasts VALUES(?,?,?,?,?,?,?)',
  (r['id'],target,r['season'],r['week'],r['captured_at'],r['kickoff'],json.dumps(r,sort_keys=True)))
 db.commit();return True
def latest_forecasts(db):
 rows=db.execute('SELECT target,payload FROM forecasts ORDER BY captured_at,id').fetchall()
 result={}
 for target,payload in rows:
  r=json.loads(payload)
  check=dict(r);saved=check.pop('id')
  if digest(check)!=saved:raise ValueError('Forecast hash mismatch')
  if parse(r['captured_at'])>=parse(r['kickoff']):raise ValueError('Non-pregame forecast in ledger')
  result[target]=r
 return result
def latest_grades(db):
 result={}
 for target,payload in db.execute('SELECT target,payload FROM grades ORDER BY graded_at,id'):
  result[target]=json.loads(payload)
 return result
def prob(loc,sd,lower,line):
 if sd==0:return float(loc>line),float(loc<line),float(loc==line)
 cdf=lambda x:statistics.NormalDist(loc,sd).cdf(x)
 over=1. if line<lower else 1-cdf(math.floor(line)+.5)
 under=0. if math.ceil(line)-1<lower else cdf(math.ceil(line)-.5)
 return over,under,max(0.,1-over-under)
def grade(db,final_games,actuals,snaps,state=STATE):
 """Grade only final games. An absent stat row is never silently treated as zero."""
 previous=latest_grades(db);added=0
 for target,f in latest_forecasts(db).items():
  if f['game_id'] not in final_games:continue
  if utcnow()<=parse(f['kickoff']):continue
  key=(f['game_id'],f['player_id']);a=actuals.get(key);snap=snaps.get((f['game_id'],f.get('pfr_id','')))
  if a is None and snap is None and not (f['availability'].startswith('Out') or f['availability'].startswith('Not playing')):continue
  played=(snap is not None and snap>0) or (a is not None and any(a.get(k,0)>0 for k in ['pass_attempts','rush_attempts','targets']))
  if not played:
   result={'status':'DNP / void','actual':None,'error':None,'z':None,'performance_eligible':False,'prop_results':[]}
  else:
   # Positive offensive snaps establishes a genuine played-zero when stats are absent.
   actual=(a or {}).get(f['metric'],0.)
   error=actual-f['location'];raw_sd=f['raw_sd'];eligible=f['workload']>0 and raw_sd>0 and f['calibrated']
   prop_results=[]
   for q in (f.get('quotes',[]) if f['workload']>0 else []):
    result='Push' if actual==q['line'] else 'Over' if actual>q['line'] else 'Under'
    valid_time=bool(q.get('quote_utc')) and parse(q['quote_utc'])<=parse(f['captured_at'])<parse(f['kickoff'])
    po,pu,pp=prob(f['location'],f['sd'],f['lower'],q['line'])
    prop_results.append(dict(book=q['book'],line=q['line'],result=result,
      probability_eligible=valid_time and result!='Push',
      brier=(po/(po+pu)-int(result=='Over'))**2 if valid_time and result!='Push' else None,
      provenance='Third-party quote; hypothetical evaluation, no bet placed'))
   result={'status':'Final','actual':actual,'error':error,'z':(actual-f['raw_mean'])/raw_sd if eligible else None,
     'performance_eligible':eligible,'covered80':f['p10']<=actual<=f['p90'],
     'covered50':abs(actual-f['location'])<=.67448975*f['sd'],'prop_results':prop_results}
  provenance={'final_game_source':final_games[f['game_id']], 'stats_source':URLS['stats.csv'],'snaps_source':URLS['snaps.csv']}
  if target in previous and previous[target]['result']==result:continue
  g={'target':target,'forecast_id':f['id'],'player':f['player'],'position':f['position'],
    'metric':f['metric'],'season':f['season'],'week':f['week'],'result':result,'source':provenance,'graded_at':iso(dt.datetime.now(UTC))}
  g['id']=digest(g);save_json(state/'grades'/f"{g['id']}.json",g)
  db.execute('INSERT OR IGNORE INTO grades VALUES(?,?,?,?,?)',(g['id'],target,f['id'],g['graded_at'],json.dumps(g,sort_keys=True)));added+=1
 db.commit();return added
def load_reconstruction(state=STATE,seed=None):
 pointer=state/'reconstruction_current.json'
 if not pointer.exists():return None
 link=json.loads(pointer.read_text());r=json.loads((state/'reconstruction.json').read_text())
 if digest({k:v for k,v in r.items() if k not in ['created_at','id']})!=r['id'] or r['id']!=link['id']:raise ValueError('Reconstruction archive hash mismatch')
 if r['kind']!='RECONSTRUCTED_NOT_LIVE' or r['through_week']!=7:raise ValueError('Invalid reconstruction scope')
 if seed and r['source_sha256']!=seed['source_sha256']:raise ValueError('Reconstruction source mismatch')
 return r

def recalibrate(db,seed,state=STATE,completed_weeks=None):
 """Shrink live standardized errors to the historical prior; versions never change past forecasts."""
 replay=load_reconstruction(state,seed); reconstructed=[r for r in replay['records'] if r['eligible']] if replay else []
 grades={k:g for k,g in latest_grades(db).items() if completed_weeks is None or g['week'] in completed_weeks};params={};live_counts={};pseudo_n=50
 for base in seed['calibration']:
  z=[max(-4.,min(4.,g['result']['z'])) for g in grades.values()
    if g['position']==base['pos'] and g['metric']==base['metric'] and g['result']['performance_eligible']]
  rz=[max(-4.,min(4.,r['z'])) for r in reconstructed if r['position']==base['pos'] and r['metric']==base['metric']]
  n=len(z);rn=.25*len(rz);mean=(pseudo_n*base['bias_z']+sum(z)+.25*sum(rz))/(pseudo_n+n+rn)
  second=(pseudo_n*(base['scale']**2+base['bias_z']**2)+sum(v*v for v in z)+.25*sum(v*v for v in rz))/(pseudo_n+n+rn)
  scale=max(.5,min(2.5,math.sqrt(max(.01,second-mean*mean))))
  params[base['key']]={'bias_z':mean,'scale':scale,'live_n':n,'reconstructed_n':len(rz),'reconstruction_effective_n':rn};live_counts[base['key']]=n
 payload={'schema':1,'training_digest':digest([replay['id'] if replay else None,sorted(g['id'] for g in grades.values())]),
  'reconstruction_id':replay['id'] if replay else None,'reconstruction_through_week':7 if replay else None,'reconstruction_weight':.25,'reconstructed_records':len(reconstructed),
  'parameters':params,'historical_prior_equivalent_n':pseudo_n,'standardized_residual_clip':4,
  'method':'Mean and dispersion recalibration only; no automatic feature/role-weight optimization',
  'latest_completed_week':max((g['week'] for g in grades.values()),default=None)}
 version=digest(payload);record=dict(payload,id=version,created_at=iso(dt.datetime.now(UTC)))
 db.execute('INSERT OR IGNORE INTO parameter_versions VALUES(?,?,?)',(version,record['created_at'],json.dumps(record,sort_keys=True)));db.commit()
 existing=json.loads(db.execute('SELECT payload FROM parameter_versions WHERE id=?',(version,)).fetchone()[0])
 save_json(state/'calibration_current.json',existing);return existing
def kickoff(game):
 return dt.datetime.fromisoformat(game['gameday']+'T'+game['gametime']).replace(tzinfo=ZoneInfo('America/New_York')).astimezone(UTC)
def normalized_stats(rows):
 out=[]
 for row in rows:
  if row['season_type']!='REG' or row['position'] not in ['QB','RB','WR','TE']:continue
  d={MAPPING.get(k,k):num(v) for k,v in row.items() if k in MAPPING or k in ['targets','receptions','receiving_yards','receiving_tds']}
  d.update(player=row['player_display_name'],player_id=row['player_id'],team=row['team'],opponent=row['opponent_team'],
    season=int(row['season']),week=int(row['week']),game_id=row['game_id'],position=row['position'])
  out.append(d)
 return out
def availability(player,roster,report,espn,game,overrides,now,news_cutoff=None):
 """Reserve status always beats absence from the weekly injury report."""
 k=norm(player['name']);rs=[r for r in roster if norm(r['full_name'])==k]
 if k=='kennethgainwell':rs=[r for r in roster if norm(r['full_name'])=='kennygainwell']
 rs=[r for r in rs if r['position']==player['pos']] or rs
 r=next((r for r in rs if r['status']=='ACT'),rs[0] if rs else {})
 evidence=[];status=None;detail='';source=URLS['roster.csv'];raw='';workload=1.
 override=overrides.get('availability',{}).get(player['name'])
 if override and override.get('season')==2026 and parse(override['effective_at'])<=now<parse(override['expires_at']):
  status=override['status'];detail=override['reason'];source=override['source']
 elif r.get('status')=='RES':
  status='Out (IR)' if r.get('status_description_abbr') in ['R01','R48'] else 'Out (reserve)';detail=r.get('status_description_abbr','Reserve list')
 elif r.get('status') in ['DEV','CUT','RET','EXE','TRT','TRC']:
  status='Not playing (roster)';detail=r['status']
 ir=next((z for z in report if z.get('gsis_id')==r.get('gsis_id')),None)
 news=next((z for z in espn if str(z.get('athlete',{}).get('id',''))==str(int(num(r.get('espn_id'))))),None)
 if news is None:news=next((z for z in espn if norm(z.get('athlete',{}).get('displayName',''))==norm(r.get('full_name',''))),None)
 if news and now-parse(news['date'])>dt.timedelta(days=8):news=None
 if news and news.get('status')!='Injured Reserve' and news_cutoff and parse(news['date'])<news_cutoff:news=None
 if ir:raw=ir.get('report_status','');detail=ir.get('report_primary_injury') or ir.get('practice_primary_injury') or detail
 if status is None:
  ns=news.get('status') if news else ''
  source=URLS['espn.json'] if news else URLS['injuries.csv']
  if ns=='Injured Reserve':status='Out (IR)'
  elif ns=='Out' or raw=='Out':
   txt=(news or {}).get('shortComment','').lower()
   status='Not playing (inactive)' if "coach's decision" in txt else 'Out (injury)'
  elif ns=='Active':status='Playing'
  elif ns=='Doubtful' or raw=='Doubtful':status='Questionable';raw='Doubtful'
  elif ns=='Questionable' or raw=='Questionable':status='Questionable'
  elif r.get('status')=='ACT':status='Playing';source=URLS['roster.csv']
  elif r.get('status')=='INA':status='Not playing (inactive)'
  else:status='Unverified'
 if news and status!='Playing':detail=news.get('shortComment','') or detail
 if not game:status='Bye';detail='No scheduled game this week'
 if status.startswith('Out') or status.startswith('Not playing') or status in ['Bye','Unverified']:workload=0.
 # Available is expected, not a fictitious guarantee of future participation.
 return dict(status=status,workload=workload,detail=detail,source=source,raw_designation=raw,
  roster_status=r.get('status','Unknown'),player_id=r.get('gsis_id',''),pfr_id=r.get('pfr_id',''),
  current_team=r.get('team',player['source_team']),practice=(ir or {}).get('practice_status',''),
  observed_at=iso(now),expected_not_guaranteed=status=='Playing')
def project(seed,stats,roster,inj,espn,games,week,params,overrides,now,config=None):
 config=config or {'prior_games':4.,'recent_weight':.5,'efficiency_trials':40.}
 before=[r for r in stats if r['week']<week];team={};defense={};gamekeys={}
 for r in before:
  team.setdefault(r['team'],{});defense.setdefault(r['opponent'],{});gamekeys.setdefault(r['team'],set()).add(r['game_id'])
  for k in ['pass_attempts','rush_attempts','targets','pass_yards','rush_yards','receiving_yards']:
   team[r['team']][k]=team[r['team']].get(k,0)+r.get(k,0);defense[r['opponent']][k]=defense[r['opponent']].get(k,0)+r.get(k,0)
 totalgames=sum(map(len,gamekeys.values()));leagueplays=sum(t.get('pass_attempts',0)+t.get('rush_attempts',0) for t in team.values())/max(1,totalgames)
 gs={t:g for g in games if int(g['week'])==week for t in [g['away_team'],g['home_team']]}
 first=min(kickoff(g) for g in games if int(g['week'])==week).astimezone(ZoneInfo('America/New_York'))
 news_cutoff=(first-dt.timedelta(days=first.weekday())).replace(hour=0,minute=0,second=0,microsecond=0).astimezone(UTC)
 rows=[]
 for p in seed['players']:
  identity=next((r for r in roster if norm(r['full_name'])==norm(p['name']) and r['position']==p['pos']),{})
  player_team=identity.get('team',p['source_team']);g=gs.get(player_team)
  av=availability(p,roster,inj,espn,g,overrides,now,news_cutoff);player_team=av['current_team'];g=gs.get(player_team)
  opp=(g['away_team'] if g['home_team']==player_team else g['home_team']) if g else 'BYE'
  ss=[r for r in before if r['player_id']==av['player_id']]
  recent=sorted(ss,key=lambda r:r['week'])[-3:];n=len(ss);weight=n/(n+config['prior_games'])
  total=lambda k:sum(r.get(k,0) for r in ss)
  last=lambda k:sum(r.get(k,0) for r in recent)/len(recent) if recent else 0.
  plays=lambda t,side:(side.get(t,{}).get('pass_attempts',0)+side.get(t,{}).get('rush_attempts',0))/max(1,len(gamekeys.get(t,set()))) if t in gamekeys else leagueplays
  pace=max(.9,min(1.1,(plays(player_team,team)+plays(opp,defense))/(2*max(1,leagueplays))))
  def mult(metric,den):
   league=sum(t.get(metric,0) for t in team.values())/max(1,sum(t.get(den,0) for t in team.values()))
   opp_rate=defense.get(opp,{}).get(metric,0)/max(1,defense.get(opp,{}).get(den,0)) if opp in gamekeys else league
   ng=len(gamekeys.get(opp,set()));return max(.85,min(1.15,1+ng/(ng+8)*(opp_rate/max(.01,league)-1)))
  match={'pass_yards':mult('pass_yards','pass_attempts'),'rush_yards':mult('rush_yards','rush_attempts'),'receiving_yards':mult('receiving_yards','targets')}
  fac=overrides.get('player_factors',{}).get(p['name'],{})
  if fac.get('week',week)!=week:fac={}
  pace*=fac.get('pace',1);role=av['workload']*fac.get('workload',1)
  uncertainty=fac.get('uncertainty',1.15 if av['status']=='Questionable' else 1.)
  components={};details={}
  for metric in METRICS[:-1]:
   den='pass_attempts' if metric in ['pass_yards','pass_tds','interceptions'] else 'rush_attempts' if metric in ['rush_yards','rush_tds'] else 'targets'
   pv=p['base'].get(metric,0);po=p['base'].get(den,0);co=total(den);l3=last(den)
   if metric=='fumbles':po=p['base']['rush_attempts']+p['base']['receptions'];co=total('rush_attempts')+total('receptions');l3=last('rush_attempts')+last('receptions')
   ct=total(metric);recent_weight=config['recent_weight'];trials=config['efficiency_trials']
   projected=(po*(1-weight)+weight*((co/n if n else 0)*(1-recent_weight)+l3*recent_weight))*pace*role
   rate=(trials*(pv/po if po else 0)+ct)/(trials+co);m=match.get(metric,1.)*fac.get(metric+'_matchup',1.)
   components[metric]=projected*rate*m;details[metric]=dict(prior=pv,prior_opp=po,total=ct,opp_total=co,recent_opp=l3)
  a=components;components['ppr']=.04*a['pass_yards']+.1*(a['rush_yards']+a['receiving_yards'])+a['receptions']+4*a['pass_tds']+6*(a['rush_tds']+a['receiving_tds'])-2*(a['interceptions']+a['fumbles'])
  distributions={}
  for metric in METRICS:
   cal=params['parameters'].get(p['pos']+'|'+metric);rawsd=p['sd'].get(metric,0)*uncertainty;mean=components[metric]
   loc=mean+rawsd*cal['bias_z'] if cal and role else mean;sd=rawsd*cal['scale'] if cal and role else 0.
   lower=0 if metric in ['receptions','pass_yards','pass_tds','rush_tds','receiving_tds','interceptions','fumbles'] else -10
   qs=[max(lower,statistics.NormalDist(loc,sd).inv_cdf(q)) if sd else loc for q in [.1,.5,.9]]
   if metric!='ppr':qs=[max(lower,math.floor(q+.5)) if sd else loc for q in qs]
   distributions[metric]=dict(raw_mean=mean,raw_sd=rawsd,location=loc,sd=sd,lower=lower,p10=qs[0],p50=qs[1],p90=qs[2],calibrated=bool(cal))
  match={k:v*fac.get(k+'_matchup',1.) for k,v in match.items()}
  rows.append(dict(name=p['name'],pos=p['pos'],team=player_team,opp=opp,n=n,availability=av,role=role,uncertainty=uncertainty,pace=pace,match=match,game=g,
    kickoff=iso(kickoff(g)) if g else None,last_week=max((r['week'] for r in ss),default=0),details=details,distributions=distributions))
 return rows,before
def sync_saved_edits(overrides,config,week):
 """Carry saved analyst-factor edits into the next same-week refresh, never override IR."""
 path=ROOT.parent/'Fantasy_2026_Weekly_Learning_Model.xlsx';lastfile=STATE/'workbook_baseline.json'
 if not path.exists() or not lastfile.exists():return
 import openpyxl
 wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
 last=json.loads(lastfile.read_text());previous={p['name']:p for p in last['players']}
 if int(wb['Method']['B3'].value)!=week:return
 for row in wb['Player inputs'].iter_rows(min_row=5,values_only=True):
  if row[0] not in previous:continue
  old=previous[row[0]];fac=overrides.setdefault('player_factors',{}).setdefault(row[0],{'week':week})
  if fac.get('week')!=week:fac.clear();fac['week']=week
  for col,key,oldval in [(6,'workload',old['role']),(7,'uncertainty',old['uncertainty']),(8,'pace',old['pace']),
      (9,'pass_yards_matchup',old['match']['pass_yards']),(10,'rush_yards_matchup',old['match']['rush_yards']),(11,'receiving_yards_matchup',old['match']['receiving_yards'])]:
   value=row[col]
   if isinstance(value,(float,int)) and oldval and abs(value-oldval)>1e-8:
    if key=='uncertainty':fac[key]=max(.5,min(3,value))
    else:fac[key]=fac.get(key,1)*max(0,min(3,value))/oldval
 for addr,key in [('B5','prior_games'),('B6','recent_weight'),('B7','efficiency_trials')]:
  v=wb['Method'][addr].value
  if isinstance(v,(int,float)) and (0<=v<=1 if key=='recent_weight' else v>0):config[key]=v
 quotes=[]
 for r in wb['Props'].iter_rows(min_row=5,max_row=44,values_only=True):
  if not r[0] or not r[1] or not isinstance(r[3],(float,int)):continue
  stamp=r[6]
  if isinstance(stamp,dt.datetime):stamp=iso(stamp.replace(tzinfo=UTC))
  elif isinstance(stamp,(int,float)):stamp=iso(dt.datetime.fromtimestamp((stamp-25569)*86400,UTC))
  else:stamp=None
  quotes.append(dict(player=r[0],metric=r[1],book=r[2],line=r[3],over_odds=r[4],under_odds=r[5],quote_utc=stamp,source=r[8],season=2026,week=week))
 save_json(ROOT/'quotes.json',quotes);save_json(ROOT/'overrides.json',overrides);save_json(ROOT/'config.json',config)
def completed_games(games,offline):
 # Scores alone do not establish finality; ESPN explicitly completed=true is required.
 final={}
 for date in sorted({g['gameday'] for g in games}):
  if date>dt.datetime.now(UTC).date().isoformat():continue
  path=CACHE/('scoreboard_'+date+'.json');url='https://site.api.espn.com/apis/site/v2/sports/football/nfl/scoreboard?dates='+date.replace('-','')
  if not offline:fetch(url,path)
  if not path.exists():continue
  data=json.loads(path.read_text());ids={str(g['espn']):g['game_id'] for g in games}
  for event in data.get('events',[]):
   if event.get('status',{}).get('type',{}).get('completed') and str(event.get('id')) in ids:final[ids[str(event['id'])]]=url
 return final
def cycle(offline=False,week=None,render=True,state=STATE):
 STATE.mkdir(parents=True,exist_ok=True);CACHE.mkdir(parents=True,exist_ok=True)
 sources={}
 if not offline:
  for filename,url in URLS.items():sources[filename]=fetch(url,CACHE/filename)
  save_json(CACHE/'manifest.json',sources)
 elif (CACHE/'manifest.json').exists():sources=json.loads((CACHE/'manifest.json').read_text())
 for filename in URLS:
  if not (CACHE/filename).exists():raise RuntimeError('Missing required source '+filename)
 seed=json.loads((ROOT/'seed.json').read_text());overrides=json.loads((ROOT/'overrides.json').read_text())
 source_file=ROOT/'seed.json'
 if len(seed['players'])!=197 or len({p['name'] for p in seed['players']})!=197:raise RuntimeError('Locked 197-player universe mismatch')
 config=json.loads((ROOT/'config.json').read_text()) if (ROOT/'config.json').exists() else {'prior_games':4.,'recent_weight':.5,'efficiency_trials':40.}
 now=dt.datetime.now(UTC);games=[g for g in readcsv(CACHE/'games.csv') if int(g['season'])==2026 and g['game_type']=='REG']
 stats=normalized_stats(readcsv(CACHE/'stats.csv'));roster=readcsv(CACHE/'roster.csv');injuries=readcsv(CACHE/'injuries.csv')
 db=db_open(state);fs=latest_forecasts(db);forecast_weeks={f['week'] for f in fs.values()};targetgames=[g for g in games if int(g['week']) in forecast_weeks]
 finals=completed_games(targetgames,offline);actuals={(r['game_id'],r['player_id']):r for r in stats}
 snaps={(r['game_id'],r['pfr_player_id']):num(r['offense_snaps']) for r in readcsv(CACHE/'snaps.csv')}
 complete_weeks={week for week in forecast_weeks if all(g['game_id'] in finals for g in games if int(g['week'])==week)}
 new_grades=grade(db,finals,actuals,snaps,state);parameters=recalibrate(db,seed,state,complete_weeks)
 future=[g for g in games if kickoff(g)>now]
 season_complete=not future
 week=week or (min(int(g['week']) for g in future) if future else max(int(g['week']) for g in games))
 # Cloud config is versioned with the engine; never read another machine's workbook.
 # sync_saved_edits is reserved for the local companion.
 latest_roster={}
 for r in sorted(roster,key=lambda r:int(num(r['week']))):
  if int(num(r['week']))<=week:latest_roster[r.get('gsis_id') or norm(r['full_name'])]=r
 roster=list(latest_roster.values())
 report=[r for r in injuries if int(num(r['week']))==week]
 espnraw=json.loads((CACHE/'espn.json').read_text());espn=[i for t in espnraw.get('injuries',[]) for i in t.get('injuries',[])]
 rows,before=project(seed,stats,roster,report,espn,games,week,parameters,overrides,now,config)
 quotes=json.loads((ROOT/'quotes.json').read_text());captured=0;late=0
 observed=min((v['observed_at'] for v in sources.values()),default=iso(now))
 for p in rows:
  if not p['game']:continue
  if parse(p['kickoff'])<=now:late+=1;continue
  if p['availability']['status']=='Unverified' or not p['availability']['player_id']:continue
  for metric,v in p['distributions'].items():
   if not v['calibrated']:continue
   record=dict(season=2026,week=week,game_id=p['game']['game_id'],player=p['name'],player_id=p['availability']['player_id'],pfr_id=p['availability']['pfr_id'],
     position=p['pos'],metric=metric,kickoff=p['kickoff'],data_through_week=week-1,inputs_observed_at=observed,
     model_version='opportunity-v2-reserve-aware-'+digest(config)[:8],parameter_version=parameters['id'],workload=p['role'],availability=p['availability']['status'],
     quotes=[q for q in quotes if q['season']==2026 and q['week']==week and q['player']==p['name'] and q['metric']==metric],**v)
   captured+=int(freeze(db,record,state))
 grades=list(latest_grades(db).values());history=[]
 by_target=latest_grades(db)
 for target,f in latest_forecasts(db).items():
  g=by_target.get(target,{}).get('result',{})
  history.append(dict(season=f['season'],week=f['week'],game_id=f['game_id'],player=f['player'],metric=f['metric'],availability=f['availability'],captured_at=f['captured_at'],kickoff=f['kickoff'],raw_mean=f['raw_mean'],p10=f['p10'],p50=f['p50'],p90=f['p90'],actual=g.get('actual'),error=g.get('error'),grade=g.get('status','Awaiting final game'),parameter_version=f['parameter_version'],forecast_id=f['id']))
 if history:
  with (state/'forecast_history.csv').open('w',newline='') as out:
   writer=csv.DictWriter(out,fieldnames=list(history[0]));writer.writeheader();writer.writerows(history)
 counts={}
 for p in rows:counts[p['availability']['status']]=counts.get(p['availability']['status'],0)+1
 report=dict(reconstruction=load_reconstruction(state,seed),generated_at=iso(now),season=2026,week=week,players=rows,actuals=before,parameters=parameters,grades=grades,
   sources=sources,status_counts=counts,config=config,season_complete=season_complete,new_frozen_metric_forecasts=captured,started_players_excluded=late,
   total_frozen_forecasts=db.execute('SELECT COUNT(*) FROM forecasts').fetchone()[0],graded_targets=len(grades),new_grades=new_grades,
   live_calibration_records=sum(g['result']['performance_eligible'] and g['week'] in complete_weeks for g in grades),source_sha256=seed['source_sha256'],completed_learning_weeks=sorted(complete_weeks),forecast_history=history)
 save_json(state/'current_run.json',report)
 from export_playerboard import export
 export(report)
 if render:subprocess.run([NODE,str(ROOT/'render.mjs'),str(state/'current_run.json')],check=True,cwd=WORKSPACE)
 return {k:report[k] for k in ['season','week','status_counts','new_frozen_metric_forecasts','started_players_excluded','total_frozen_forecasts','graded_targets','new_grades','live_calibration_records']}

def main():
 parser=argparse.ArgumentParser();parser.add_argument('command',choices=['cycle','status']);parser.add_argument('--offline',action='store_true');parser.add_argument('--week',type=int);parser.add_argument('--no-render',action='store_true');args=parser.parse_args()
 STATE.mkdir(parents=True,exist_ok=True)
 with (STATE/'run.lock').open('a') as lock:
  fcntl.flock(lock,fcntl.LOCK_EX)
  if args.command=='status':
   d=json.loads((STATE/'current_run.json').read_text());print(json.dumps({k:d[k] for k in ['generated_at','season','week','status_counts','total_frozen_forecasts','graded_targets','live_calibration_records']},indent=2));return
  print(json.dumps(cycle(args.offline,args.week,not args.no_render),indent=2))
if __name__=='__main__':main()
